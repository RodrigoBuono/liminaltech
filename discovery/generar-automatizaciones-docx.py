"""
Script para convertir el documento de automatizaciones de Aurea a Word (.docx)
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_automatizaciones_docx():
    md_path = Path("/home/user/liminaltech/discovery/automatizaciones-aurea.md")
    docx_path = Path("/home/user/liminaltech/discovery/docx/automatizaciones-proceso-interno-aurea.docx")
    docx_path.parent.mkdir(exist_ok=True)

    text = md_path.read_text(encoding="utf-8")
    doc = Document()

    # Estilos base
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Saltar lineas vacias y separadores
        if stripped == "" or stripped == "---":
            i += 1
            continue

        # Titulo principal (# )
        if stripped.startswith("# ") and not stripped.startswith("## ") and not stripped.startswith("### "):
            title = stripped[2:]
            p = doc.add_heading(title, level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Subtitulo (## )
        if stripped.startswith("## ") and not stripped.startswith("### "):
            subtitle = stripped[3:]
            doc.add_heading(subtitle, level=1)
            i += 1
            continue

        # Sub-subtitulo (### )
        if stripped.startswith("### "):
            subsubtitle = stripped[4:]
            p = doc.add_heading(subsubtitle, level=2)
            i += 1
            continue

        # Bold text with **
        if "**" in stripped:
            p = doc.add_paragraph()
            import re
            parts = re.split(r"(\*\*.*?\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.size = Pt(11)
                    r.font.name = "Calibri"
                else:
                    r = p.add_run(part)
                    r.font.size = Pt(11)
                    r.font.name = "Calibri"
            i += 1
            continue

        # Texto normal
        if stripped:
            p = doc.add_paragraph(stripped)
            for run in p.runs:
                run.font.size = Pt(11)
                run.font.name = "Calibri"

        i += 1

    doc.save(str(docx_path))
    print(f"Generado: {docx_path}")


if __name__ == "__main__":
    create_automatizaciones_docx()
